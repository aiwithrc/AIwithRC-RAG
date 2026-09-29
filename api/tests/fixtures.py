"""Small test documents built in code, so no binary fixtures live in the repo."""

import io
import zlib


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF, one page per string, one text line per '\\n'-separated paragraph."""
    objs: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        lines, y = [], 760
        for para in text.split("\n"):
            safe = para.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            lines.append(f"BT /F1 11 Tf 50 {y} Td ({safe}) Tj ET")
            y -= 16
        stream = zlib.compress("\n".join(lines).encode("latin-1"))
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * i} 0 R >>".encode()
        )
        objs.append(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_blank_pdf(n_pages: int = 2) -> bytes:
    """Pages with no text layer, like a scanned document."""
    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(n_pages):
        w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """paragraphs: (style, text) with style like 'Heading 1' or 'Normal'."""
    import docx

    d = docx.Document()
    for style, text in paragraphs:
        if style.startswith("Heading"):
            d.add_heading(text, level=int(style.split()[-1]))
        else:
            d.add_paragraph(text)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeEmbedder:
    """Deterministic bag-of-words hashing embedder: fast, no model download, similar texts ≈ similar vectors."""

    model_name = "fake"
    dim = 64

    def _vec(self, text: str) -> list[float]:
        import math
        import re

        v = [0.0] * self.dim
        for w in re.findall(r"\w+", text.lower()):
            v[zlib.crc32(w.encode()) % self.dim] += 1.0
        norm = math.sqrt(sum(x * x for x in v)) or 1.0
        return [x / norm for x in v]

    def embed_passages(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)
