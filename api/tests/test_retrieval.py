"""Retrieval, citations, confidence and index QA."""

import pytest

from app.jobs import queue
from app.providers.chat import ThinkFilter
from app.rag import cite, rerank
from app.rag.chunk import context_header, with_context
from app.rag.parse import parse
from app.rag.retrieve import Passage, retrieve, rrf
from tests.fixtures import make_pdf
from tests.test_ingest import MSA, upload


class KeywordReranker:
    """Fake cross-encoder: score = shared words between query and text (logit-ish scale)."""

    def scores(self, query, texts):
        import re

        q = {w for w in re.findall(r"\w+", query.lower()) if len(w) > 3}
        # Mimics ms-marco-MiniLM: about -11 for unrelated text, higher per shared word.
        return [-11.0 + 5.0 * sum(w in t.lower() for w in q) for t in texts]


@pytest.fixture(autouse=True)
def fake_reranker():
    rerank.set_reranker(KeywordReranker())
    yield
    rerank.set_reranker(None)


def P(n: int, text: str, score: float = 0.9, page: int | None = 1, section: str | None = None) -> Passage:
    return Passage(
        chunk_id=f"c{n}", document_id="d1", filename="Acme_MSA.pdf", page=page, section=section, text=text,
        ordinal=n - 1, doc_chunks=10, score=score,
    )


# ---- RRF ----

def test_rrf_rewards_agreement_between_rankings():
    fused = rrf([["a", "b", "c"], ["c", "a", "d"]])
    order = [i for i, _ in fused]
    assert order[0] == "a"  # 1/61 + 1/62 beats c's 1/63 + 1/61
    assert order[1] == "c"
    assert set(order) == {"a", "b", "c", "d"}
    assert fused[0][1] == pytest.approx(1 / 61 + 1 / 62)


def test_rrf_single_list_keeps_order():
    assert [i for i, _ in rrf([["x", "y", "z"]])] == ["x", "y", "z"]


# ---- citations ----

def test_renumber_drops_invalid_and_orders_by_first_use():
    text, used = cite.renumber("Notice is 60 days [3]. Breach is 30 days [1, 3]. Nonsense [9].", 3)
    assert text == "Notice is 60 days [1]. Breach is 30 days [2][1]. Nonsense."
    assert used == [3, 1]


def test_renumber_handles_other_marker_styles():
    text, used = cite.renumber("A 【2】 and B [^1] and C [1-2].", 2)
    assert used == [2, 1]
    assert text == "A [1] and B [2] and C [2][1]."


def test_strip_think():
    assert cite.strip_think("<think>hmm let me see</think>\n\nThe answer [1].") == "The answer [1]."
    assert cite.strip_think("reasoning without opening tag</think>Answer.") == "Answer."


def test_think_filter_streaming_across_chunk_boundaries():
    f = ThinkFilter()
    parts = ["<th", "ink>secret ", "stuff</thi", "nk>Hello", " world <", "b>"]
    out = "".join(f.feed(p) for p in parts) + f.flush()
    assert out == "Hello world <b>"


def test_highlight_picks_supporting_sentence():
    chunk = (
        "11.1 Term. This Agreement starts on the Effective Date. "
        "11.2 Termination for Convenience. Either party may terminate this Agreement for any reason upon sixty (60) days' "
        "prior written notice to the other party. Fees remain payable."
    )
    before, hit, after = cite.highlight(chunk, "Either party can terminate with 60 days' written notice [1].")
    assert hit.startswith("Either party may terminate")
    assert "Effective Date" in before and "Fees remain payable" in after
    assert (before + hit + after).replace("  ", " ").strip() == chunk


def test_build_citations_and_confidence():
    passages = [P(1, "Invoices are due within forty-five days.", 0.91), P(2, "Termination needs sixty days notice.", 0.8)]
    text, cites = cite.build_citations("<think>x</think>Termination needs **60 days** notice [2].", passages)
    assert text == "Termination needs **60 days** notice [1]."
    assert [(c.n, c.chunk_id, c.chunk, c.total) for c in cites] == [(1, "c2", 2, 10)]
    assert cites[0].hit == "Termination needs sixty days notice."
    assert cite.confidence(passages, cites) == "high"
    assert cite.confidence([P(1, "x", 0.6)], cites) == "low"  # best relevance < 0.75
    assert cite.confidence(passages, []) == "low"  # no valid citation


# ---- retrieval against a real index ----

def test_retrieve_hybrid_rerank_and_scores(client, owner):
    kb = owner["default_kb_id"]
    upload(client, kb, ("Acme_MSA.pdf", MSA, "application/pdf"),
           ("notes.txt", b"The cafeteria opens at nine. Parking is free on weekends.", "text/plain"))
    queue.process_all()
    from app.db import SessionLocal

    with SessionLocal() as db:
        ps = retrieve(db, kb, "How many days notice to terminate the agreement?", top_k=3)
        assert ps[0].section == "§11.2 Termination for Convenience" and ps[0].page == 2
        assert 0 <= ps[-1].score <= ps[0].score <= 1
        assert ps[0].score >= 0.75
        assert ps[0].location() == "Page 2 · §11.2 Termination for Convenience"
        # Without rerank, scores are cosine similarities and still ordered.
        plain = retrieve(db, kb, "cafeteria parking weekends", top_k=2, rerank=False, hybrid=False)
        assert plain[0].filename == "notes.txt" and plain[0].score >= plain[-1].score
        assert retrieve(db, "no-such-kb", "anything") == []


# ---- parsing improvements found by QA ----

def test_docx_visual_headings_bold_and_larger_font(tmp_path):
    import docx
    from docx.shared import Pt

    d = docx.Document()
    def para(text, size, bold=False):
        r = d.add_paragraph().add_run(text)
        r.font.size, r.bold = Pt(size), bold
    para("JANE DOE", 16, bold=True)
    para("Singapore | jane@example.com", 10)
    para("Education", 11, bold=True)
    para("National University of Singapore, B.Eng in Computer Engineering, 2012. Graduated with honours.", 10)
    para("Experience", 11, bold=True)
    para("Staff Engineer - Acme | 2020 - Present", 10.5)
    para("Led the platform team and shipped the billing system. Mentored six engineers.", 10)
    p = tmp_path / "abc123"  # stored without extension, like real uploads
    d.save(p)
    blocks = parse(p, "Jane_Doe_CV.docx")
    by_text = {b.text: b.section for b in blocks}
    assert by_text["National University of Singapore, B.Eng in Computer Engineering, 2012. Graduated with honours."] == "JANE DOE › Education"
    assert by_text["Led the platform team and shipped the billing system. Mentored six engineers."] == (
        "JANE DOE › Experience › Staff Engineer - Acme | 2020 - Present"
    )


def test_context_header():
    assert context_header("Acme_MSA_2024.pdf", "§11.2 Termination") == "Acme MSA 2024 › §11.2 Termination"
    assert with_context("notes.txt", None, "Body") == "notes\nBody"


# ---- index QA ----

def test_check_reports_healthy_and_broken_documents(client, owner):
    from app.db import SessionLocal
    from app.models import Chunk
    from app.rag import store
    from app.rag.health import check_all
    from tests.fixtures import make_blank_pdf

    kb = owner["default_kb_id"]
    upload(client, kb, ("Acme_MSA.pdf", MSA, "application/pdf"), ("scan.pdf", make_blank_pdf(), "application/pdf"),
           ("long.pdf", make_pdf([" ".join(f"Clause {i} covers topic {i} in detail." for i in range(120))]), "application/pdf"))
    queue.process_all()
    with SessionLocal() as db:
        reports = {r.filename: r for r in check_all(db)}
        assert reports["Acme_MSA.pdf"].ok, reports["Acme_MSA.pdf"].problems
        assert reports["Acme_MSA.pdf"].self_retrieval[0] == reports["Acme_MSA.pdf"].self_retrieval[1] > 0
        assert not reports["scan.pdf"].ok and "No text layer" in reports["scan.pdf"].problems[0]
        # Break one vector: QA must notice.
        victim = db.query(Chunk).filter(Chunk.text.contains("Clause 1 ")).first()
        store.collection(kb).delete(ids=[victim.id])
        broken = {r.filename: r for r in check_all(db)}["long.pdf"]
        assert any("missing a vector" in p for p in broken.problems)


def test_relevance_calibration_matches_measurements():
    from app.rag.cite import HIGH_CONFIDENCE, MIN_RELEVANCE
    from app.rag.rerank import relevance

    # Measured ms-marco-MiniLM logits: unrelated pairs ~ -11, answering passages -0.2 .. +6.6.
    for unrelated in (-11.4, -10.8):
        assert relevance(unrelated) < MIN_RELEVANCE
    for answering in (-0.22, 0.17, 1.73, 3.10, 6.63):
        assert relevance(answering) >= HIGH_CONFIDENCE
    assert relevance(-5.0) == pytest.approx(0.5)
